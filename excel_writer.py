"""
抽出した物件データを Excel に書き出す。

- 既存の賃料査定書（.xlsx）を渡した場合：
  「物件名」と「〜賃料」の見出しがあるシート（＝競合案件シート）を探し、
  見出しの名前で列を割り当てて、最後の行の次に同じ書式で追記する。
  表紙・スタッキングなど他のシートや既存の行には手を触れない。
- 渡さない場合：同じ見出しの新しい表を作る。
"""

import copy
import io
import re
import unicodedata

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# 項目キー → 見出しとして認める名前（空白・改行・全角半角は無視して照合）
HEADER_ALIASES = {
    'name': ['物件名', '建物名'],
    'age_str': ['築年数', '築年月', '築年'],
    'area': ['㎡数', '面積', '専有面積', 'm2数'],
    'layout': ['間取', '間取り'],
    'address': ['住所', '所在地'],
    'station': ['駅', '最寄駅', '最寄り駅'],
    'walk': ['徒歩', '徒歩分'],
    'structure': ['構造'],
    'total_rent': ['合計賃料', '総賃料', '賃料合計'],
    'deposit': ['敷金'],
    'key_money': ['礼金'],
    'ad': ['広告料', 'AD', '広告費'],
    'unit_price': ['㎡単価', '平米単価', 'm2単価'],
    'notes': ['特筆事項', '備考'],
    'sepa': ['セパ', 'バストイレ別'],
    'ev': ['EV', 'エレベーター'],
    'al': ['AL', 'オートロック'],
    'basin': ['洗面台', '独立洗面台'],
    'laundry': ['室内洗濯機', '室内洗濯'],
    'dressing': ['脱衣所'],
    'bath_dry': ['浴室乾燥'],
    'reheat': ['追炊き', '追焚き'],
    'kitchen': ['キッチン'],
    'storage': ['収納'],
    'shampoo': ['シャンドレ', 'シャンプードレッサー'],
    'shoe_box': ['シューズボックス', 'シューズBOX'],
    'delivery': ['宅配ボックス', '宅配BOX'],
    'ac': ['エアコン'],
    'washlet': ['ウォシュレット', '温水洗浄便座'],
}

# 新規作成時の見出し（第一住建の競合案件シートと同じ並び＋追加設備）
DEFAULT_HEADERS = [
    ('name', '物件名'), ('age_str', '築年数'), ('area', '㎡数'), ('layout', '間取'),
    ('address', '住所'), ('station', '駅'), ('walk', '徒歩'), ('structure', '構造'),
    ('total_rent', '合計賃料'), ('deposit', '敷金'), ('key_money', '礼金'), ('ad', '広告料'),
    ('unit_price', '㎡単価'), ('notes', '特筆事項'),
    ('sepa', 'セパ'), ('ev', 'ＥＶ'), ('al', 'ＡＬ'), ('basin', '洗面台'),
    ('laundry', '室内\n洗濯機'), ('dressing', '脱衣所'), ('kitchen', 'キッチン'),
    ('storage', '収納'), ('shampoo', 'シャンドレ'), ('washlet', 'ｳｫｼｭﾚｯﾄ'),
    ('bath_dry', '浴室乾燥'), ('reheat', '追炊き'), ('delivery', '宅配\nボックス'), ('ac', 'エアコン'),
]


def _norm(v):
    if v is None:
        return ''
    s = unicodedata.normalize('NFKC', str(v))
    return re.sub(r'\s', '', s).upper()


_ALIAS_LOOKUP = {_norm(a): k for k, names in HEADER_ALIASES.items() for a in names}


def find_target(wb):
    """(シート, 見出し行, {項目キー: 列番号}) を返す。見つからなければ None。"""
    found = []
    for ws in wb.worksheets:
        for r in range(1, min(ws.max_row, 20) + 1):
            row = {c: _norm(ws.cell(r, c).value) for c in range(1, min(ws.max_column, 60) + 1)}
            vals = set(row.values())
            if '物件名' in vals and any('賃料' in v for v in vals):
                cols = {}
                for c, v in row.items():
                    k = _ALIAS_LOOKUP.get(v)
                    if k and k not in cols:
                        cols[k] = c
                found.append((ws, r, cols))
                break
    if not found:
        return None
    # 「競合」を含むシートを優先
    found.sort(key=lambda x: ('競合' not in x[0].title))
    return found[0]


def _last_data_row(ws, header_row, name_col):
    last = header_row
    for r in range(header_row + 1, ws.max_row + 1):
        if ws.cell(r, name_col).value not in (None, ''):
            last = r
    return last


def _value(prop, key, row, cols):
    if key == 'unit_price':
        if 'total_rent' in cols and 'area' in cols:
            t = get_column_letter(cols['total_rent'])
            a = get_column_letter(cols['area'])
            return f'=IF(OR({a}{row}="",{a}{row}=0),"",{t}{row}/{a}{row})'
        return ''
    v = prop.get(key, '')
    return '' if v is None else v


def append_to_template(template_bytes, properties):
    """既存の査定書に追記。戻り値は (xlsxのbytes, シート名, 追記した行の範囲)。"""
    wb = openpyxl.load_workbook(io.BytesIO(template_bytes))
    target = find_target(wb)
    if not target:
        raise ValueError('Excelに「物件名」「合計賃料」などの見出しがあるシートが見つかりませんでした。')
    ws, header_row, cols = target
    name_col = cols['name']
    last = _last_data_row(ws, header_row, name_col)
    style_row = last if last > header_row else header_row + 1
    max_col = max(cols.values())
    start = last + 1
    for i, prop in enumerate(properties):
        r = start + i
        h = ws.row_dimensions[style_row].height
        if h:
            ws.row_dimensions[r].height = h
        for c in range(1, max_col + 1):
            ws.cell(r, c)._style = copy.copy(ws.cell(style_row, c)._style)
        for key, c in cols.items():
            ws.cell(r, c).value = _value(prop, key, r, cols)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue(), ws.title, (start, start + len(properties) - 1)


def build_new(properties):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '競合案件'
    thin = Side(style='thin', color='999999')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    hdr_fill = PatternFill('solid', fgColor='1F4E78')
    cols = {}
    ws.row_dimensions[3].height = 32
    for ci, (key, label) in enumerate(DEFAULT_HEADERS, 1):
        c = ws.cell(3, ci, label)
        c.fill = hdr_fill
        c.font = Font(name='Meiryo UI', bold=True, color='FFFFFF', size=10)
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        c.border = border
        cols[key] = ci
    widths = {'name': 30, 'address': 30, 'notes': 40, 'age_str': 11, 'ad': 16, 'kitchen': 12, 'storage': 10}
    for key, ci in cols.items():
        ws.column_dimensions[get_column_letter(ci)].width = widths.get(key, 8 if ci > 14 else 10)
    for i, prop in enumerate(properties):
        r = 4 + i
        ws.row_dimensions[r].height = 30
        for key, ci in cols.items():
            c = ws.cell(r, ci, _value(prop, key, r, cols))
            c.font = Font(name='Meiryo UI', size=10)
            c.border = border
            c.alignment = Alignment(vertical='center', wrap_text=True,
                                    horizontal='center' if ci > 14 or key in ('walk', 'layout', 'structure') else None)
            if key in ('total_rent', 'deposit', 'key_money'):
                c.number_format = '#,##0'
            if key == 'unit_price':
                c.number_format = '#,##0'
    ws.freeze_panes = 'B4'
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue(), ws.title, (4, 3 + len(properties))
