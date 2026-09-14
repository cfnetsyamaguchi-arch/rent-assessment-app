#!/usr/bin/env python3
"""
賃料査定書 PDF自動入力アプリ (Streamlit版)

ローカル起動:
    pip install -r requirements.txt
    streamlit run app.py
"""

import io
import re
import copy
import json
from datetime import date

import pandas as pd
import streamlit as st
import pdfplumber
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment

CURRENT_YEAR = date.today().year

# ── PDF Parser ──────────────────────────────────────────────────────────────
WAREKI = {'明治': 1868, '大正': 1912, '昭和': 1926, '平成': 1989, '令和': 2019}


def wareki_to_year(s):
    for era, base in WAREKI.items():
        m = re.search(era + r'(\d+)年', s)
        if m:
            return base + int(m.group(1)) - 1
    m = re.search(r'(\d{4})年', s)
    return int(m.group(1)) if m else None


def parse_num(s):
    if not s:
        return None
    s = re.sub(r'[,，]', '', str(s))
    m = re.search(r'\d+', s)
    return int(m.group()) if m else None


def has_eq(text, *patterns):
    return '〇' if any(re.search(p, text) for p in patterns) else '×'


# 敷金・礼金の直後に続きうる別項目のラベル（ここで打ち切る）
NEXT_LABEL = r'(礼金|敷金|保証金|敷引|償却|解約|管理費|共益費|水道|駐車|更新|仲介|A\s*D|AD)'


def grab_field(label, text, width=20):
    """label の直後の値を、次の項目ラベルの手前まで取り出す。"""
    m = re.search(label + r'[：:\s]*([^\n]{0,%d})' % width, text)
    if not m:
        return ''
    return re.split(NEXT_LABEL, m.group(1))[0].strip()


def to_yen(raw, rent):
    """敷金・礼金の記載を金額(円)に正規化する。

    「1ヶ月」「1.5ヶ月」のような月数表記は、総賃料ではなく家賃(rent)を基準に換算する。
    家賃が取得できていない場合は換算せず、原文をそのまま返す。
    """
    if raw is None:
        return ''
    s = re.sub(r'[,，\s]', '', str(raw))
    if not s:
        return ''
    if re.search(r'(なし|無し|不要|ゼロ|^[-－ー―]+$)', s):
        return 0

    # 月数表記（1ヶ月 / 1.5ヵ月 / 2カ月 …）→ 家賃 × 月数
    m = re.search(r'([\d.]+)\s*[ヶヵケカか箇]?月', s)
    if m:
        months = float(m.group(1))
        if months == 0:
            return 0
        return int(round(rent * months)) if rent else str(raw).strip()

    # 万円表記
    m = re.search(r'([\d.]+)\s*万', s)
    if m:
        return int(float(m.group(1)) * 10000)

    # 数値のみ。家賃より明らかに小さい端数は月数表記とみなす（例:「敷金 1」＝1ヶ月）
    m = re.search(r'\d+(?:\.\d+)?', s)
    if m:
        v = float(m.group())
        if v == 0:
            return 0
        if v < 100 and rent:
            return int(round(rent * v))
        return int(v)

    return ''


def extract_text(file_bytes):
    with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
        pages = [p.extract_text() for p in pdf.pages if p.extract_text()]
    text = '\n'.join(pages)
    if not text.strip():
        raise ValueError('スキャンPDF（画像）のため自動抽出できません。手動で入力してください。')
    return text


def parse_pdf(text):
    d = {}

    # 物件名
    name = ''
    m = re.search(r'物件名\s+([^\n]+)', text)
    if m:
        raw = m.group(1).strip()
        # 備考テキストが混入（※で始まる）場合は次の行から取得
        if raw.startswith('※') or ('！' in raw and len(raw) > 15):
            bm = re.search(r'\n([^\n※☆\d「」\s]{2,25})\s+(\d{3})\s+種', text)
            if bm:
                name = bm.group(1).strip()
        else:
            name = raw

    # 建物名フォールバック
    if not name:
        m = re.search(r'建物名\s+([^\n\s種]+)', text)
        if m:
            name = m.group(1).strip()

    # 号室（数字のみ有効）
    room = ''
    m = re.search(r'号室名?\s+(\d{2,4})', text)
    if m:
        room = m.group(1) + '号室'
    else:
        bm = re.search(r'[^\n]{2,25}\s+(\d{3})\s+種', text)
        if bm:
            room = bm.group(1) + '号室'
        else:
            m2 = re.search(r'^(\d{3})\s+[\d.]+\s+即入', text, re.MULTILINE)
            if m2:
                room = m2.group(1) + '号室'

    if room and room.replace('号室', '') not in name:
        d['name'] = (name + ' ' + room).strip()
    else:
        d['name'] = name

    # 築年数（「築年月 平成3年3月」「築年 1997年07月」両対応）
    m = re.search(r'築年(?:月)?\s+([^\n]+)', text)
    if m:
        year = wareki_to_year(m.group(1))
        if year:
            d['age'] = CURRENT_YEAR - year
            d['age_str'] = f'{year}年'  # 和暦・西暦どちらの記載でも西暦に統一

    # 面積
    m = re.search(r'(?:専有面積|使用部分)[^\d]*([\d.]+)\s*㎡', text)
    if m:
        d['area'] = float(m.group(1))

    # 間取（標準パターン以外は空白）
    m = re.search(r'間取(?:タイプ|内訳|り)?\s+([^\n]+)', text)
    if m:
        lm = re.search(r'(ワンルーム|[1-9][LSDK]+)', m.group(1))
        if lm:
            d['layout'] = lm.group(1)
    if 'layout' not in d:
        lm = re.search(r'(ワンルーム|[1-9][LSDK]+)', text)
        if lm:
            d['layout'] = lm.group(1)

    # 住所（複数パターン対応）
    addr_found = False
    for pat in [
        r'所在地\s*\n?\s*([^\n]+(?:丁目|番地?|号)[^\n]*)',
        r'在\s*地?\s+([^\n]+(?:丁目|番地?)[^\n]*)',
        r'在\s+([^\n]+(?:丁目|番地?)[^\n]*)',
    ]:
        m = re.search(pat, text)
        if m:
            addr = re.sub(r'〒[\d-]+', '', m.group(1)).strip()
            am = re.search(r'(?:大阪|東京|京都|兵庫|奈良|滋賀|和歌山)[^\n]{3,}', addr)
            d['address'] = am.group().strip() if am else addr
            addr_found = True
            break
    # 大阪市/府で始まる行を直接探す（縦組みレイアウト対応）
    if not addr_found:
        m = re.search(r'^((?:大阪府|大阪市|東京都|京都府)[^\n]+(?:丁目|番地?|号)[^\n]*)', text, re.MULTILINE)
        if m:
            d['address'] = m.group(1).strip()

    # 駅・徒歩（最寄り）
    hits = re.findall(r'[「「]([^「」「」\n]{1,20}?)[」」]\s*徒歩\s*(\d+)\s*分', text)
    hits += re.findall(r'([^\s「」「」\n]{2,20}駅)\s*徒歩\s*(\d+)\s*分', text)
    if hits:
        closest = min(hits, key=lambda x: int(x[1]))
        st_name = closest[0].strip()
        sm = re.search(r'(\S+駅|\S+)$', st_name)
        d['station'] = sm.group(1) if sm else st_name
        d['walk'] = int(closest[1])

    # 構造
    m = re.search(r'(SRC造|RC造|鉄骨造|木造|鉄筋コンクリート造)', text)
    if m:
        mapping = {'SRC造': 'SRC', 'RC造': 'RC', '鉄筋コンクリート造': 'RC',
                   '鉄骨造': '鉄骨', '木造': '木造'}
        d['structure'] = mapping.get(m.group(1), m.group(1))

    # 賃料（「賃料 50,000円」と「50,000 円\n賃料」両対応）
    rent = 0
    m = re.search(r'賃料\s+([\d,]+)\s*円', text)
    if m:
        rent = parse_num(m.group(1)) or 0
    else:
        m = re.search(r'([\d,]+)\s*円\s*\n賃料', text)
        if m:
            rent = parse_num(m.group(1)) or 0
    d['rent'] = rent

    # 管理費・共益費・水道代
    mgmt = 0
    m = re.search(r'(?:管理費等?|共益費(?:・管理費)?)[：:\s]+([\d,]+)\s*円', text)
    if m:
        mgmt = parse_num(m.group(1)) or 0
    water = 0
    m = re.search(r'水道料金[：:\s]+([\d,]+)\s*円', text)
    if m:
        water = parse_num(m.group(1)) or 0
    d['mgmt'] = mgmt
    d['total_rent'] = rent + mgmt + water

    # 敷金・礼金（月数表記は「家賃」基準で金額に換算。総賃料は使わない）
    d['deposit'] = to_yen(grab_field('敷金', text), rent)
    d['key_money'] = to_yen(grab_field('礼金', text), rent)

    # 広告料（「A D 2ヶ月」「AD 150%」等、%優先）
    m = re.search(r'A\s*D\s+(\d+)\s*(%|[ヶヵ]月|ヶ月|ヵ月|月)?', text)
    if m:
        val = m.group(1)
        unit = m.group(2) or ''
        if unit and '%' not in unit and 'ヶ月' not in unit:
            unit = unit.replace('ヵ月', 'ヶ月').replace('月', 'ヶ月')
        d['ad'] = val + unit

    # 設備（〇 / ×）
    d['ev'] = has_eq(text, 'エレベーター', 'ＥＶ')
    d['al'] = has_eq(text, 'オートロック', 'ＡＬ')
    d['laundry'] = has_eq(text, r'室内洗濯[パ機]', '室内洗濯機置場')
    d['dressing'] = has_eq(text, '独立洗面', '脱衣所')
    d['bath_dry'] = has_eq(text, '浴室乾燥')
    d['reheat'] = has_eq(text, '追炊き', '追い焚き')
    d['kitchen'] = has_eq(text, 'IHコンロ', 'ガスコンロ', 'IHキッチン', 'システムキッチン')
    d['shampoo'] = has_eq(text, 'シャンドレ', 'シャンプードレッサー')
    d['shoe_box'] = has_eq(text, 'シューズボックス', 'シューズＢＯＸ')
    d['delivery'] = has_eq(text, '宅配ボックス', '宅配ＢＯＸ')
    d['ac'] = has_eq(text, 'エアコン')
    d['washlet'] = has_eq(text, '温水洗浄便座', 'ウォシュレット')

    return d


# ── Excel Generator ─────────────────────────────────────────────────────────
YELLOW = PatternFill('solid', fgColor='FFFF00')


def build_excel(properties, template_bytes=None):
    if template_bytes:
        wb = openpyxl.load_workbook(io.BytesIO(template_bytes))
        ws = wb.active
        next_row = 5
        for r in range(5, ws.max_row + 2):
            if all(ws.cell(row=r, column=c).value is None for c in range(1, 27)):
                next_row = r
                break
    else:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = '競合案件'
        _write_headers(ws)
        next_row = 4

    for i, prop in enumerate(properties):
        _write_row(ws, next_row + i, prop, template_bytes is not None)

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return out.read()


def _write_headers(ws):
    headers = ['物件名', '築年数', '㎡数', '間取', '住所', '駅', '徒歩', '構造',
               '合計賃料', '敷金', '礼金', '広告料', '㎡単価', '特筆事項',
               'ＥＶ', 'ＡＬ', '室内\n洗濯機', '脱衣所', '浴室乾燥', '追炊き',
               'キッチン', 'シャンドレ', 'シューズボックス', '宅配ボックス',
               'エアコン', 'ｳｫｼｭﾚｯﾄ']
    hdr_fill = PatternFill('solid', fgColor='4472C4')
    hdr_font = Font(bold=True, color='FFFFFF')
    ws.row_dimensions[3].height = 30
    for ci, h in enumerate(headers, 1):
        c = ws.cell(row=3, column=ci, value=h)
        c.fill = hdr_fill
        c.font = hdr_font
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)


def _write_row(ws, row_num, prop, copy_style):
    WHITE = PatternFill('solid', fgColor='FFFFFF')
    if copy_style:
        for ci in range(1, 27):
            ref = ws.cell(row=4, column=ci)
            cell = ws.cell(row=row_num, column=ci)
            if ref.has_style:
                cell.font = copy.copy(ref.font)
                cell.border = copy.copy(ref.border)
                cell.fill = WHITE
                cell.number_format = ref.number_format
                cell.alignment = copy.copy(ref.alignment)
    else:
        for ci in range(1, 27):
            ws.cell(row=row_num, column=ci).fill = YELLOW

    vals = [
        prop.get('name', ''),
        prop.get('age_str', prop.get('age', '')),
        prop.get('area', ''),
        prop.get('layout', ''),
        prop.get('address', ''),
        prop.get('station', ''),
        prop.get('walk', ''),
        prop.get('structure', ''),
        prop.get('total_rent', ''),
        prop.get('deposit', ''),
        prop.get('key_money', ''),
        prop.get('ad', ''),
        f'=I{row_num}/C{row_num}',
        '',
        prop.get('ev', ''), prop.get('al', ''), prop.get('laundry', ''),
        prop.get('dressing', ''), prop.get('bath_dry', ''), prop.get('reheat', ''),
        prop.get('kitchen', ''), prop.get('shampoo', ''), prop.get('shoe_box', ''),
        prop.get('delivery', ''), prop.get('ac', ''), prop.get('washlet', ''),
    ]
    for ci, v in enumerate(vals, 1):
        ws.cell(row=row_num, column=ci).value = v


# ── 表示用カラム定義 ────────────────────────────────────────────────────────
COLUMNS = [
    ('name', '物件名'), ('age_str', '築年数'), ('area', '㎡数'), ('layout', '間取'),
    ('address', '住所'), ('station', '駅'), ('walk', '徒歩'), ('structure', '構造'),
    ('rent', '家賃'), ('total_rent', '合計賃料'), ('deposit', '敷金'),
    ('key_money', '礼金'), ('ad', '広告料'),
    ('ev', 'ＥＶ'), ('al', 'ＡＬ'), ('laundry', '室内洗濯'), ('dressing', '脱衣所'),
    ('bath_dry', '浴室乾燥'), ('reheat', '追炊き'), ('kitchen', 'キッチン'),
    ('shampoo', 'シャンドレ'), ('shoe_box', 'シューズBOX'), ('delivery', '宅配BOX'),
    ('ac', 'エアコン'), ('washlet', 'ウォシュレット'),
]
EQ_KEYS = ['ev', 'al', 'laundry', 'dressing', 'bath_dry', 'reheat', 'kitchen',
           'shampoo', 'shoe_box', 'delivery', 'ac', 'washlet']
# 数値として扱う項目（Excelでそのまま計算・並べ替えできるようにする）
INT_KEYS = ['walk', 'rent', 'total_rent', 'deposit', 'key_money']
FLOAT_KEYS = ['area']


def to_dataframe(records):
    rows = []
    for r in records:
        row = {}
        for key, label in COLUMNS:
            v = r.get(key, '')
            if key in EQ_KEYS and v not in ('〇', '×'):
                v = '×'
            row[label] = '' if v is None else v
        rows.append(row)
    df = pd.DataFrame(rows, columns=[label for _, label in COLUMNS])

    for key, label in COLUMNS:
        if key in INT_KEYS or key in FLOAT_KEYS:
            # float64 + NaN にすると、値なしのセルが空欄として表示される
            df[label] = pd.to_numeric(df[label], errors='coerce').astype('float64')
        else:
            df[label] = df[label].astype('string')
    return df


def from_dataframe(df):
    records = []
    for _, row in df.iterrows():
        rec = {}
        for key, label in COLUMNS:
            v = row.get(label, '')
            if v is None or (not isinstance(v, str) and pd.isna(v)):
                rec[key] = ''
            elif key in INT_KEYS:
                rec[key] = int(v)
            elif key in FLOAT_KEYS:
                rec[key] = float(v)
            else:
                rec[key] = v
        records.append(rec)
    return records


# ── Streamlit UI ────────────────────────────────────────────────────────────
st.set_page_config(page_title='賃料査定書 自動入力', page_icon='🏢', layout='wide')

st.markdown("""
<style>
  .block-container { padding-top: 2rem; max-width: 1500px; }
  h1 { font-size: 1.6rem !important; }
</style>
""", unsafe_allow_html=True)

st.title('🏢 賃料査定書 自動入力アプリ')
st.caption('物件PDFをアップロード → 内容を確認・修正 → Excelでダウンロード')

if 'records' not in st.session_state:
    st.session_state.records = []
if 'errors' not in st.session_state:
    st.session_state.errors = []

col1, col2 = st.columns([2, 1])
with col1:
    pdf_files = st.file_uploader(
        '📄 物件PDF（複数選択可）', type=['pdf'], accept_multiple_files=True)
with col2:
    tmpl_file = st.file_uploader(
        '📊 既存Excel（任意・テンプレートを引き継ぎます）', type=['xlsx'])

if st.button('📄 PDFから抽出', type='primary', disabled=not pdf_files):
    records, errors = [], []
    prog = st.progress(0.0)
    for i, f in enumerate(pdf_files):
        try:
            data = parse_pdf(extract_text(f.getvalue()))
            data['filename'] = f.name
            records.append(data)
        except Exception as e:
            errors.append(f'{f.name}: {e}')
        prog.progress((i + 1) / len(pdf_files))
    prog.empty()
    st.session_state.records = records
    st.session_state.errors = errors

for err in st.session_state.errors:
    st.warning('⚠️ ' + err)

if st.session_state.records:
    st.success(f'✅ {len(st.session_state.records)} 件を抽出しました。内容を確認・修正してください。')

    df = to_dataframe(st.session_state.records)
    col_config = {}
    for key, label in COLUMNS:
        if key in EQ_KEYS:
            col_config[label] = st.column_config.SelectboxColumn(
                label, options=['〇', '×'], width='small')
        elif key in INT_KEYS:
            col_config[label] = st.column_config.NumberColumn(
                label, format='%d', step=1)
        elif key in FLOAT_KEYS:
            col_config[label] = st.column_config.NumberColumn(label, format='%.2f')
        elif key in ('name', 'address'):
            col_config[label] = st.column_config.TextColumn(label, width='medium')

    edited = st.data_editor(
        df, column_config=col_config, num_rows='dynamic',
        width='stretch', hide_index=True, key='editor')

    records = from_dataframe(edited)
    excel_bytes = build_excel(records, tmpl_file.getvalue() if tmpl_file else None)

    st.download_button(
        '⬇ Excelをダウンロード',
        data=excel_bytes,
        file_name=f'賃料査定書_{date.today().isoformat()}.xlsx',
        mime='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        type='primary')

    if st.button('🗑 クリア'):
        st.session_state.records = []
        st.session_state.errors = []
        st.rerun()
else:
    st.info('PDFをアップロードして「PDFから抽出」を押してください。')

st.divider()
st.caption(
    '※ 敷金・礼金は「◯ヶ月」表記の場合、**家賃（管理費・水道料金を含まない金額）×月数** で円に換算しています。'
    'Excelには「家賃」列は出力されません（確認用の表示です）。'
)
st.caption('※ 築年数は和暦の記載でも西暦（例：1991年）に統一して表示します。')
st.caption('※ 画像スキャンのPDFは文字が取り出せないため自動抽出できません。その場合は表に手入力してください。')
